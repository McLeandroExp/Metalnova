"""Edita la v0.3 existente: agrega D5-D8, quita notas de tablas y aplica estilo de bibliografía."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

BASE = Path(__file__).resolve().parent
DOC = BASE / "_working_v0.3.docx"
CAP = BASE.parent / "capturas"


def add_caption(doc, ref, num, title, page_break=False):
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"Figura {num}")
    r.bold = True
    r2 = p.add_run(". ")
    r2.bold = True
    r3 = p.add_run(title)
    r3.italic = True
    p.paragraph_format.page_break_before = page_break
    ref._p.addprevious(p._p)
    return p


def add_image(doc, ref, filename, width=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(CAP / filename), width=Inches(width))
    ref._p.addprevious(p._p)
    return p


def add_note(doc, ref, text):
    p = doc.add_paragraph(style="Note")
    r = p.add_run("Nota. ")
    r.italic = True
    p.add_run(text)
    ref._p.addprevious(p._p)
    return p


def add_list_entry(doc, after_p, label, title):
    new_p = deepcopy(after_p._p)
    after_p._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    par = Paragraph(new_p, after_p._parent)
    if len(par.runs) >= 3:
        par.runs[0].text = label
        par.runs[1].text = "\t" + title
        par.runs[2].text = "\t—"
    return par


def main():
    doc = Document(DOC)

    # 1. Actualizar intro del Anexo D.
    for p in doc.paragraphs:
        if p.text.startswith("Las figuras D1–D4 reproducen"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            r = p.add_run(
                "Las figuras D1–D4 reproducen las cuatro páginas principales. Las figuras D5–D8 "
                "reproducen las páginas de apoyo: detalle de máquina, detalle de causa, tooltip y "
                "validación T-604."
            )
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            p.paragraph_format.first_line_indent = Cm(1.27)
            p.paragraph_format.line_spacing = 2.0
            p.paragraph_format.space_after = Pt(0)
            break

    # 2. Insertar D5-D8 antes de Anexo E.
    ref = None
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.startswith("Anexo E."):
            ref = p
            break
    if ref is None:
        raise RuntimeError("No se encontró Anexo E")

    capturas = [
        ("D5", "Detalle máquina", "detalle-maquina.png", True),
        ("D6", "Detalle causa", "detalle-causa.png", False),
        ("D7", "Tooltip máquina", "tooltip-maquina.png", True),
        ("D8", "Modelo validado (T-604)", "modelo-validado-t604.png", False),
    ]
    for num, title, filename, br in capturas:
        add_caption(doc, ref, num, title, page_break=br)
        add_image(doc, ref, filename, width=6.0)
        add_note(doc, ref, "Captura de Power BI Desktop, 2026-10-23.")

    # 3. Agregar D5-D8 a la lista de figuras.
    d4 = None
    for p in doc.paragraphs:
        if p.style.name == "ListEntry" and p.text.startswith("D4\t"):
            d4 = p
            break
    if d4 is not None:
        for label, title in [
            ("D5", "Detalle máquina"),
            ("D6", "Detalle causa"),
            ("D7", "Tooltip máquina"),
            ("D8", "Modelo validado (T-604)"),
        ]:
            d4 = add_list_entry(doc, d4, label, title)

    # 4. Quitar notas de tablas y la nota de paginación de las listas.
    remove_prefixes = (
        "Nota. La paginación de estas listas",
        "Nota. Elaboración propia a partir de la documentación del prototipo",
        "Nota. El diccionario completo",
        "Nota. El detalle está en docs/verification",
        "Nota. Los comandos se ejecutan",
        "Nota. Términos usados",
    )
    for p in list(doc.paragraphs):
        if any(p.text.startswith(pref) for pref in remove_prefixes):
            p._p.getparent().remove(p._p)

    # 5. Estilo de bibliografía para las referencias.
    styles = doc.styles
    try:
        bib = styles["Bibliography"]
    except KeyError:
        bib = styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
        bib.base_style = styles["Normal"]
    bib.font.name = "Times New Roman"
    bib.font.size = Pt(12)
    bib.paragraph_format.line_spacing = 2.0
    bib.paragraph_format.space_after = Pt(0)
    bib.paragraph_format.left_indent = Cm(1.27)
    bib.paragraph_format.first_line_indent = Cm(-1.27)
    bib.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rpr = bib.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")

    in_refs = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.startswith("8. Referencias"):
            in_refs = True
            continue
        if p.style.name == "Heading 1" and p.text.startswith("Anexo A."):
            break
        if in_refs and p.text.strip():
            p.style = bib
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

    doc.save(DOC)
    print(DOC)


if __name__ == "__main__":
    main()
